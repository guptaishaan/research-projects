#!/usr/bin/env python3
"""Generate student assignment .docx files (3 per notebook) for each project.

Assignment 1 - Introduction and Literature Review
Assignment 2 - Preprocessing and Methodology
Assignment 3 - Results and Discussion

Run: python tools/build_assignments.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------------------
# Small formatting helpers
# ---------------------------------------------------------------------------

def new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return doc


def title(doc, project, assignment):
    p = doc.add_paragraph()
    r = p.add_run(project)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    h = doc.add_paragraph()
    hr = h.add_run(assignment)
    hr.bold = True
    hr.font.size = Pt(16)
    doc.add_paragraph()  # spacer


def section(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12.5)
    return p


def sub(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph(text)
    if italic:
        p.runs[0].italic = True
    return p


def questions(doc, items):
    for i, q in enumerate(items, 1):
        doc.add_paragraph(f"{i}. {q}")


def paper(doc, citation, why, qs):
    sub(doc, citation)
    body(doc, why, italic=True)
    questions(doc, qs)
    doc.add_paragraph()


def instructions_block(doc, lines):
    for ln in lines:
        doc.add_paragraph(ln, style="List Bullet")


# ---------------------------------------------------------------------------
# Assignment builders (shared structure, per-notebook content passed in)
# ---------------------------------------------------------------------------

def build_a1(project, intro, papers, closing):
    doc = new_doc()
    title(doc, project, "Assignment 1 - Introduction and Literature Review")
    section(doc, "Goal")
    body(doc, intro)
    doc.add_paragraph()
    section(doc, "How to do this assignment")
    instructions_block(doc, [
        "Read each paper below (the abstract, introduction, and conclusion are enough to answer the questions).",
        "Answer the questions in your own words - two to four sentences each.",
        "Keep your answers; we will reuse them when we write the Introduction and Related Work sections of our own paper.",
    ])
    doc.add_paragraph()
    section(doc, "Papers to read")
    for c, w, qs in papers:
        paper(doc, c, w, qs)
    section(doc, "For our paper")
    body(doc, closing)
    return doc


def build_a2(project, notebook, intro, groups, closing):
    doc = new_doc()
    title(doc, project, "Assignment 2 - Preprocessing and Methodology")
    section(doc, "Goal")
    body(doc, intro)
    doc.add_paragraph()
    section(doc, "How to do this assignment")
    instructions_block(doc, [
        f"Open the notebook: {notebook}",
        "Go to the sections named below and read them, running or re-reading each cell as you go.",
        "Answer the questions in your own words. Where a figure is mentioned, describe what you actually see.",
        "Your answers become the raw material for the Data and Methods sections of our paper.",
    ])
    doc.add_paragraph()
    for heading, blurb, qs in groups:
        section(doc, heading)
        if blurb:
            body(doc, blurb)
        questions(doc, qs)
        doc.add_paragraph()
    section(doc, "For our paper")
    body(doc, closing)
    return doc


def build_a3(project, intro, metrics_primer, result_qs, limitations_q):
    doc = new_doc()
    title(doc, project, "Assignment 3 - Results and Discussion")
    section(doc, "Goal")
    body(doc, intro)
    doc.add_paragraph()
    section(doc, "First, the scoreboard: what the metrics mean")
    for name, meaning in metrics_primer:
        p = doc.add_paragraph()
        r = p.add_run(f"{name}: ")
        r.bold = True
        p.add_run(meaning)
    doc.add_paragraph()
    section(doc, "Reading the results")
    questions(doc, result_qs)
    doc.add_paragraph()
    section(doc, "Limitations and future directions")
    body(doc, "Every honest study says what it cannot yet prove. Answer this in a short paragraph:")
    questions(doc, [limitations_q])
    return doc


if __name__ == "__main__":
    from assignments_content import PROJECTS

    for key, p in PROJECTS.items():
        outdir = os.path.join(REPO, p["folder"], "assignments")
        os.makedirs(outdir, exist_ok=True)
        name = p["project"]

        d1 = build_a1(name, p["a1_intro"], p["a1_papers"], p["a1_closing"])
        d1.save(os.path.join(outdir, "Assignment 1 - Introduction and Literature Review.docx"))

        d2 = build_a2(name, p["notebook"], p["a2_intro"], p["a2_groups"], p["a2_closing"])
        d2.save(os.path.join(outdir, "Assignment 2 - Preprocessing and Methodology.docx"))

        d3 = build_a3(name, p["a3_intro"], p["a3_metrics"], p["a3_results"], p["a3_limitations"])
        d3.save(os.path.join(outdir, "Assignment 3 - Results and Discussion.docx"))

        print("wrote 3 files ->", outdir)
